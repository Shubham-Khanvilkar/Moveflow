from docx import Document
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path
src = Path("MOVE_IN_SYNC_MASTER_PRODUCTION_SPEC_V13_COMPLETE.docx")
out = Path("MOVE_IN_SYNC_MASTER_PRODUCTION_SPEC_V14_UNLIMITED_DETAILED_MASTER.docx")
if not src.exists(): print("V13 not found"); exit(1)
doc = Document(src)
for sec in doc.sections:
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
def pb(): doc.add_page_break()
def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def p(t): doc.add_paragraph(t)
def b(t): doc.add_paragraph(t, style="List Bullet")
def tbl(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers): t.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
print("Setup done")
pb()
h1("V14 - DETAILED IMPLEMENTATION BIBLE")
p("This continuation removes any practical length restriction.")
h1("103. Product-Wide Interaction Contract")
tbl(["Rule","Required behavior"],[("Primary action","One visually dominant primary action per context."),("Secondary action","Available but visually subordinate."),("Destructive action","Requires explicit confirmation and reason."),("State change","Immediately reflects resulting state."),("Error","Specific, actionable, human-readable."),("Unauthorized","No sensitive data exposed."),("Slow operation","Use background job instead of freezing."),("Unsaved changes","Warn before navigation."),("Search","Scope-aware and human-readable."),("Code/ID","Human code is primary; UUID is secondary.")])
h1("104. Application Shell")
for x in ["Left navigation can collapse.","Top bar: company/scope, search, notifications, profile.","Main content: tabs, tables, maps, drawers, wizards.","Breadcrumbs for context.","Context selectors do not silently change scope."]: b(x)
h1("105. Context Bar"); p("Company / Site / Process / Shift / Role must be visible.")
h1("106. Global Search"); b("Employee name/ID, Trip Code, Booking Code, Vehicle, Driver, Route, Document, Incident, Invoice.")
