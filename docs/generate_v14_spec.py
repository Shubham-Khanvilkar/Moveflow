from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

src = Path("MOVE_IN_SYNC_MASTER_PRODUCTION_SPEC_V13_COMPLETE.docx")
out = Path("MOVE_IN_SYNC_MASTER_PRODUCTION_SPEC_V14_UNLIMITED_DETAILED_MASTER.docx")

doc = Document(src)

for sec in doc.sections:
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Heading 1"].font.size = Pt(17)
styles["Heading 2"].font.size = Pt(13)
styles["Heading 3"].font.size = Pt(11)

def pb(): doc.add_page_break()
def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)
def p(t): doc.add_paragraph(t)
def b(t): doc.add_paragraph(t, style="List Bullet")
def n(t): doc.add_paragraph(t, style="List Number")
def code(t):
    q = doc.add_paragraph()
    r = q.add_run(t)
    r.font.name = "Courier New"
    r.font.size = Pt(8)
    return q
def tbl(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t

pb()
h1("V14 — DETAILED IMPLEMENTATION BIBLE")
p("This continuation deliberately removes any practical length restriction. Every section is implementation guidance.")
p("The application should be easy for a first-time user to understand even though the platform underneath it is large.")

h1("103. Product-Wide Interaction Contract")
tbl(["Rule","Required behavior"], [
("Primary action","One visually dominant primary action per context."),
("Secondary action","Available but visually subordinate."),
("Destructive action","Requires explicit confirmation and reason where configured."),
("State change","Immediately reflects resulting state or processing status."),
("Error","Specific, actionable, human-readable."),
("Unauthorized","No sensitive data is exposed; show clear access state."),
("Slow operation","Use progress/background job instead of freezing the page."),
("Unsaved changes","Warn before navigation when loss is possible."),
("Refresh","Preserve filters and context where practical."),
("Search","Scope-aware and human-readable."),
("Code/ID","Human code is primary; UUID is secondary/technical."),
])

h1("104. Application Shell")
h2("104.1 Desktop shell")
b("Left navigation can collapse.")
b("Top bar contains company/scope context, global search, notifications, help and profile.")
b("Main content supports tabs, tables, maps, drawers and wizards.")
b("Breadcrumbs identify the current context.")
b("Context selectors do not silently change data scope without a visible indicator.")

h2("104.2 Mobile shell")
b("Bottom navigation for high-frequency tasks.")
b("One-handed interaction.")
b("Large status and primary action.")
b("Persistent connectivity indicator where operationally important.")
b("Emergency action remains available to authorized users.")

h1("105. Context Bar")
code("Company: ABC Limited\nSite: BLR01 - Whitefield\nProcess: P01 - Customer Support\nShift: 18:30 - 03:30\nRole: TRANSPORT ADMIN")
p("A user with multiple sites/processes can change context deliberately. Switching context must refresh data and permissions.")

h1("106. Global Search Experience")
b("Employee name and Employee ID.")
b("Trip Code, Booking Code.")
b("Vehicle registration and Vehicle Code.")
b("Driver name and Driver Code.")
b("Route Code, Document Code, Incident Code, Invoice Code.")
p("Search results are filtered server-side by effective access.")

h1("107. Notification Center")
tbl(["Category","Examples","Actions"], [
("Operations","Driver unavailable, trip delayed","Open trip / dispatch"),
("Safety","SOS, unsafe pickup","Open incident / control room"),
("Compliance","Licence expiring","Open document"),
("Billing","Invoice ready, reconciliation mismatch","Open invoice"),
("System","Integration failed","Open integration"),
("Reports","Export ready","Download"),
])

h1("108. Dashboard Design Rules")
b("A dashboard is a work surface, not a wall of decorative cards.")
b("Every KPI must be calculated from real data.")
b("Clicking a KPI should reveal the underlying authorized records.")
b("A zero must mean zero, not API failure.")
b("A loading state must not look like an empty state.")
b("Filters at dashboard level must visibly affect all widgets that honor them.")

h1("109. Super Admin Complete Workspace")
p("SUPER_ADMIN is a complete functional owner-level administrator.")
tbl(["Area","Super Admin actions"], [
("Platform Control","All"),("Owner Management","All"),("Companies","Create/edit/activate/suspend/archive/restore"),
("Users/Roles/Permissions","All"),("Employees","All"),("Bookings","All"),("Trips/Dispatch","All"),
("Fleet","All"),("QR","All"),("GPS/Maps","All"),("Safety","All"),("Documents","All"),
("Reports","All"),("Internal Billing","All"),("External Billing","All"),
("Integrations","All"),("Security","All"),("Audit","All"),("System Health","All"),
])

h1("112. Transport Admin — Complete Operational Workspace")
tbl(["Module","Primary actions"], [
("Employees","Search/add/onboard/edit/offboard/import/export"),
("Bookings","Create/edit/cancel/bulk/manage"),
("Trips","Find/create/view/assign/unassign/reassign/split/merge"),
("Passenger Operations","Add/remove/move/reassign/no-show"),
("Dispatch","Live assignment and exception handling"),
("Drivers","View/create/edit/assign/state/documents"),
("Vehicles","View/create/edit/assign/maintenance/documents"),
("Vehicle QR","Generate/regenerate/revoke/history"),
("Routes","View/create/edit/publish/version"),
("Schedules","View/create/edit/publish"),
("Roster","View/manage/import"),
("Safety","Incidents/SOS/female transport"),
("Vendors","View/manage/compliance/performance"),
("Documents","Upload/verify/reject/download/history"),
("Reports","Operational reports and permitted exports"),
])

h1("113. Trip Domain — Detailed Rules")
h2("113.1 Trip record")
tbl(["Field","Requirement"], [
("Trip UUID","Technical primary key"),("Trip Code","Mandatory human-readable code"),
("Date","Operational local date"),("Shift","Shift reference"),
("Company","Mandatory"),("Site","Mandatory"),
("Driver assignment","Versioned"),("Vehicle assignment","Versioned"),
("Status","State machine"),("Version","Concurrency control"),
])

h2("113.2 Trip code generation")
code("TRP-YYYYMMDD-NNNNNN\nExample: TRP-20260905-001274")

h1("114. Passenger Container Model")
p("A trip contains passenger records. Passenger operations must be independent from the driver/vehicle assignment.")
code("TRIP\n  +-- TRIP_PASSENGER A\n  +-- TRIP_PASSENGER B\n  +-- TRIP_PASSENGER C\n  +-- TRIP_PASSENGER D")

h1("115. Passenger Move — Full Workflow")
code("Select passenger -> Click Move Passenger -> Load eligible destination trips\n-> Filter by date/time/company/site/process/policy\n-> Check capacity/pickup window/route/safety\n-> Show impact -> Capture reason -> Confirm -> Transaction\n-> Source trip updated -> Destination trip updated -> Events written")

h1("121. Booking Domain — Detailed")
h2("121.1 Booking statuses")
code("DRAFT\nREQUESTED\nCONFIRMED\nWAITLISTED\nASSIGNED\nCANCELLED\nEXPIRED\nCOMPLETED\nNO_SHOW")

h1("127. Driver State Machine")
code("OFFLINE -> ACTIVE -> AVAILABLE -> ON_TRIP / BREAK / SLEEPY / EMERGENCY\nOFFLINE -> ACTIVE -> SUSPENDED")

h1("130. Document Platform — Full Specification")
h2("130.1 Document categories")
b("Driver licence, permit, background verification, training")
b("Vehicle registration, insurance, fitness, permit")
b("Vendor contract, compliance certificate")
b("Company contract, site permit")
b("Incident evidence, invoice attachment")

h1("131. Safety Module — Detailed")
b("Active SOS, Female transport exceptions")
b("Unsafe pickup warnings, Route deviations")
b("Driver fatigue, Guard incidents")
b("Open incidents, Respo
