from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re, os, textwrap, shutil, json, math

out = Path("/mnt/data/MOVE_IN_SYNC_MASTER_PRODUCTION_SPEC_V13_COMPLETE.docx")

doc = Document()

# Page setup
for sec in doc.sections:
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

# Base styles
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)

for name, size, bold in [
    ("Title", 24, True),
    ("Heading 1", 17, True),
    ("Heading 2", 13, True),
    ("Heading 3", 11, True),
]:
    st = styles[name]
    st.font.name = "Aptos Display" if name in ("Title","Heading 1") else "Aptos"
    st.font.size = Pt(size)
    st.font.bold = bold

# helper styles
if "Small" not in styles:
    s = styles.add_style("Small", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Aptos"
    s.font.size = Pt(8)
if "Callout" not in styles:
    s = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Aptos"
    s.font.size = Pt(9.5)
    s.paragraph_format.left_indent = Inches(0.2)
    s.paragraph_format.right_indent = Inches(0.2)
if "CodeBlock" not in styles:
    s = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Courier New"
    s.font.size = Pt(8)

def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Move In Sync V13  |  Page ")
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = "1"
    r.append(t)
    fld.append(r)
    footer._p.append(fld)

for sec in doc.sections:
    add_page_number(sec)

def page_break():
    doc.add_page_break()

def title(text):
    p = doc.add_paragraph()
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p

def h1(text):
    doc.add_heading(text, level=1)

def h2(text):
    doc.add_heading(text, level=2)

def h3(text):
    doc.add_heading(text, level=3)

def para(text, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def bullet(text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    doc.add_paragraph(text, style=style)

def numbered(text):
    doc.add_paragraph(text, style="List Number")

def code(text):
    p = doc.add_paragraph(style="CodeBlock")
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run().add_break()
        p.add_run(line)
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, x in enumerate(headers):
        hdr[i].text = x
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = str(x)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t

# Cover
title("MOVE IN SYNC")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MASTER PRODUCTION SPECIFICATION V13")
r.bold = True
r.font.size = Pt(16)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Complete Enterprise Employee Transportation Platform")
r.bold = True
r.font.size = Pt(14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Owner hierarchy • RBAC • transport operations • passenger-level dispatch • fleet • QR • GPS • safety • documents • billing • reporting • integrations • UX • Supabase").italic = True
page_break()

h1("Document Control")
table(
    ["Item","Value"],
    [
        ("Document", "Move In Sync Master Production Specification"),
        ("Version", "V13"),
        ("Purpose", "Single consolidated implementation contract for Mimo 2.5 Free"),
        ("Status", "Master build specification"),
        ("Page count rule", "No artificial page limit"),
        ("Authentication", "Supabase Auth"),
        ("Database", "Supabase Postgres"),
        ("Primary maps", "Google Maps Platform through provider abstraction"),
        ("Primary owner authority", "SAAS_OWNER / MOVEINSYNC_OWNER / SUPER_ADMIN"),
        ("Implementation principle", "Build complete functionality before cosmetics"),
    ]
)
para("This document consolidates the original RBAC master requirements and all later functional expansions into one implementation baseline. It is intentionally longer than a UI brief. It defines what the system must do, who may do it, where they may do it, what must be stored, what users see, what APIs must exist, and how production readiness is proven.")
page_break()

h1("How Mimo Must Use This Specification")
para("Mimo 2.5 Free must treat this document as the authoritative product contract. Do not compress requirements simply to reduce the number of files, routes, pages or components. Build as many screens, drawers, wizards, detail pages, reports and workflows as required to make the product usable.")
bl([
    "Never replace a real workflow with mock data.",
    "Never use hard-coded KPI values in production.",
    "Never rely on frontend checks as the security boundary.",
    "Never confuse a role with a permission or scope.",
    "Never give Manager, Finance, HRMS or Employee QR-generation authority merely because they can access another operational screen.",
    "Never treat a trip as immutable after driver assignment; trips are actively managed objects.",
    "Never treat a four-passenger trip as an indivisible unit; passengers can move independently.",
    "Never bury document upload inside unrelated settings.",
    "Never mix Move In Sync SaaS billing with customer/vendor transport billing.",
    "Never expose server-side secrets or unrestricted map keys to browsers or mobile apps.",
    "Do not declare a feature complete until its UI, API, persistence, authorization, validation, audit and acceptance test all exist."
])
page_break()

# Core architecture
h1("1. Product Vision and Scope")
para("Move In Sync is an enterprise, multi-tenant employee transportation management platform. It must support the complete journey from customer onboarding and employee transport eligibility through booking, roster, route planning, driver/vehicle assignment, dispatch, live GPS, passenger boarding, safety, trip completion, billing, reconciliation and reporting.")
h2("1.1 Primary actors")
table(
    ["Actor","Primary outcome"],
    [
        ("SAAS_OWNER","Owns the SaaS business and governs owner identities and the entire platform."),
        ("MOVEINSYNC_OWNER","Owns Move In Sync operations and has full functional platform access."),
        ("SUPER_ADMIN","Full-functional owner-level administrator across all modules and tenants."),
        ("PLATFORM_ADMIN","Day-to-day platform administration and support."),
        ("TRANSPORT_HEAD","Customer transport governance and performance."),
        ("TRANSPORT_ADMIN","Full transport operations in assigned scope; employee master authority."),
        ("TRANSPORT_COORDINATOR","Dispatch, schedule and live operations."),
        ("MANAGER / TEAM_LEADER","Team-level transport actions and approvals."),
        ("EMPLOYEE","Own transport and permitted self-service."),
        ("VENDOR_ADMIN","Vendor organization and assigned customer scope."),
        ("VENDOR_DISPATCHER","Vendor-side trip and driver dispatch."),
        ("DRIVER","Assigned duty and trip execution."),
        ("GUARD","Site-side vehicle/driver/trip verification."),
        ("FINANCE","Billing, invoice and reconciliation workflows."),
        ("SECURITY / AUDITOR","Access, security and evidence.")
    ]
)
h2("1.2 Core operating principles")
bl([
    "The platform is multi-tenant and scope-aware.",
    "Company, site, process and team are first-class dimensions.",
    "Owner-level accounts retain complete functional authority.",
    "Transport Admin owns transport employee onboarding/offboarding.",
    "Trip and passenger operations can be modified independently when operationally safe.",
    "Safety and compliance always outrank optimization.",
    "Historical operational events are retained; corrections are represented as new events or controlled adjustments.",
    "Documents have lifecycle, versioning and verification.",
    "Reports and exports obey the same authorization boundary as source data.",
])
page_break()

h1("2. High-Level System Architecture")
code("""Browser / Mobile App
        |
        v
Supabase Auth / Session
        |
        v
Application API + Domain Services
        |
        +--> Authorization + Scope Engine
        +--> Booking / Trip / Dispatch
        +--> Fleet / Driver / Vendor
        +--> QR / GPS / Safety
        +--> Documents / Reports / Billing
        |
        v
Supabase Postgres + RLS
        |
        +--> transactional tables
        +--> audit logs
        +--> outbox events
        +--> reporting/read models

External services through adapters
        |
        +--> Google Maps Platform
        +--> Email / SMS / Push / WhatsApp
        +--> HRMS / ERP / SSO
        +--> GPS providers
""")
para("Use a modular monolith or clearly bounded services initially. Do not force a microservice architecture prematurely. Domain boundaries must be explicit so high-volume components can later scale independently.")
h2("2.1 Mandatory architecture layers")
table(
    ["Layer","Responsibility"],
    [
        ("Presentation","Role-aware desktop/mobile UI, forms, maps, tables, notifications."),
        ("API","Versioned authenticated endpoints; idempotency; validation; error contracts."),
        ("Authorization","Role, permission, scope, resource and policy evaluation."),
        ("Domain","Booking, trip, fleet, safety, billing and document business rules."),
        ("Persistence","Postgres transactional state and RLS."),
        ("Async","Outbox, workers, report jobs, notifications, webhooks."),
        ("Integrations","Provider adapters; no vendor-specific logic leaking into domain."),
        ("Observability","Logs, metrics, traces, health, job state and audit."),
    ]
)
page_break()

h1("3. Ownership and Governance")
h2("3.1 Owner hierarchy")
code("""SAAS_OWNER
   |
   +--> MOVEINSYNC_OWNER
   |
   +--> SUPER_ADMIN
             |
             +--> platform/customer/admin roles
             +--> transport roles
             +--> finance roles
             +--> security/audit roles
             +--> vendor/driver/guard roles
""")
para("SAAS_OWNER is the highest governance principal. MOVEINSYNC_OWNER is the business owner of the Move In Sync platform operation. SUPER_ADMIN is a complete functional administrator. Owner hierarchy controls who may create and manage owner-level identities; it does not make SUPER_ADMIN a reduced operational role.")
h2("3.2 Owner creation")
table(
    ["Actor","Create SAAS_OWNER","Create MOVEINSYNC_OWNER","Create SUPER_ADMIN","Functional platform access"],
    [
        ("SAAS_OWNER","Yes","Yes","Yes","All"),
        ("MOVEINSYNC_OWNER","No","No","Yes","All"),
        ("SUPER_ADMIN","No","No","Yes when owner policy permits","All functional"),
    ]
)
h2("3.3 Owner Management workflow")
bl([
    "Open Owner Management.",
    "Choose owner tier.",
    "Enter identity and verified email.",
    "Require strong authentication/MFA before activation.",
    "Review owner permissions.",
    "Set status and activation date.",
    "Confirm with reason.",
    "Create owner account and audit event.",
    "Send invitation through trusted server-side Supabase Auth flow.",
    "Show activation state and audit history."
])
h2("3.4 SUPER_ADMIN all-functional rule")
para("SUPER_ADMIN is seeded with ALL_FUNCTIONAL_PERMISSIONS. An ordinary customer role cannot remove SUPER_ADMIN functional authority. Owner-level governance controls remain separate from business-function permissions. System-integrity protections such as immutable audit storage, credential secrecy and database protection must not be interpreted as business restrictions.")
page_break()

h1("4. Security Domains")
table(
    ["Domain","Meaning","Isolation requirement"],
    [
        ("PLATFORM_INTERNAL","Move In Sync platform employees","Cross-platform authority according to role."),
        ("CUSTOMER_INTERNAL","Customer company users","Only customer-authorized tenant scope."),
        ("VENDOR_EXTERNAL","Transport vendor users","Vendor identity plus assigned customer/site/process scope."),
        ("DRIVER_EXTERNAL","Driver users","Own driver, vehicle/duty and assigned trips."),
        ("GUARD_EXTERNAL","Site security users","Assigned site and operational verification scope."),
    ]
)
para("A role name is never sufficient to establish authorization. The effective identity is a combination of security domain, organization, role assignment, permission set and scope.")
page_break()

h1("5. RBAC Model")
h2("5.1 Three concepts")
table(
    ["Concept","Example","Meaning"],
    [
        ("Role","TRANSPORT_ADMIN","A named bundle of default responsibilities."),
        ("Permission","trip.reassign","One action the user may perform."),
        ("Scope","ABC / BLR01 / P01","Where the action is allowed."),
    ]
)
h2("5.2 Effective access")
code("""effective_access =
 authenticated
 AND active_account
 AND active_role
 AND permission_enabled
 AND allowed_company_scope
 AND allowed_site_scope
 AND allowed_process_scope
 AND allowed_team_scope
 AND resource_in_scope
 AND policy_allows
 AND system_deny = false
""")
h2("5.3 Server-side enforcement")
para("The server must resolve identity and scope from trusted records. Never trust role_id, company_id, site_id, process_id or user_id supplied by the browser as proof of authority.")
h2("5.4 Delegation")
para("Delegation is permission-level, scoped and auditable. A Transport Admin may delegate vehicle_qr.create to a Coordinator within an allowed scope without turning that Coordinator into a Transport Admin.")
page_break()

h1("6. Complete Role Catalogue")
roles = [
("PLATFORM","SAAS_OWNER","Ultimate SaaS ownership and governance"),
("PLATFORM","MOVEINSYNC_OWNER","Move In Sync business ownership"),
("PLATFORM","SUPER_ADMIN","Complete functional platform administrator"),
("PLATFORM","PLATFORM_ADMIN","Platform administration"),
("PLATFORM","PLATFORM_OPERATIONS","Cross-tenant operational monitoring"),
("PLATFORM","PLATFORM_SUPPORT","Customer support and diagnostics"),
("PLATFORM","PLATFORM_SECURITY_ADMIN","IAM, sessions, security, audit"),
("PLATFORM","PLATFORM_AUDITOR","Read-only platform audit"),
("PLATFORM","PLATFORM_FINANCE_ADMIN","SaaS billing and platform revenue"),
("PLATFORM","PLATFORM_FINANCE_VIEWER","Read-only SaaS billing"),
("PLATFORM","PLATFORM_BILLING_ADMIN","Customer subscription administration"),
("PLATFORM","PLATFORM_INTEGRATION_ADMIN","APIs, integrations, webhooks"),
("PLATFORM","PLATFORM_COMPLIANCE_ADMIN","Platform compliance"),
("PLATFORM","PLATFORM_CUSTOMER_SUCCESS","Customer success operations"),
("PLATFORM","PLATFORM_NOC_OPERATOR","System and operational monitoring"),
("PLATFORM","PLATFORM_DATA_ADMIN","Reporting/data administration"),
("CUSTOMER","COMPANY_ADMIN","Company-wide customer administration"),
("CUSTOMER","COMPANY_SUB_ADMIN","Delegated company administration"),
("CUSTOMER","DIRECTOR","Executive multi-site/process visibility"),
("CUSTOMER","TRANSPORT_HEAD","Transport governance"),
("CUSTOMER","TRANSPORT_ADMIN","Full transport operations"),
("CUSTOMER","TRANSPORT_SUB_ADMIN","Delegated transport operations"),
("CUSTOMER","TRANSPORT_COORDINATOR","Scheduling and dispatch"),
("CUSTOMER","DISPATCHER","Live dispatch"),
("CUSTOMER","CONTROL_ROOM_OPERATOR","Control-room operations"),
("CUSTOMER","ROSTER_ADMIN","Roster/shift administration"),
("CUSTOMER","ROSTER_PLANNER","Planning and roster preparation"),
("CUSTOMER","ROUTE_ADMIN","Routes, stops, geofences"),
("CUSTOMER","FLEET_MANAGER","Fleet operations"),
("CUSTOMER","SAFETY_ADMIN","Safety governance"),
("CUSTOMER","FEMALE_TRANSPORT_ADMIN","Female transport policy operations"),
("CUSTOMER","EMERGENCY_RESPONSE_OFFICER","Emergency response"),
("CUSTOMER","INCIDENT_MANAGER","Incident case management"),
("CUSTOMER","VENDOR_MANAGER","Vendor management"),
("CUSTOMER","VENDOR_COMPLIANCE_MANAGER","Vendor document/compliance"),
("CUSTOMER","FINANCE_ADMIN","External transport finance"),
("CUSTOMER","FINANCE_APPROVER","Finance approval"),
("CUSTOMER","FINANCE_VIEWER","Read-only finance"),
("CUSTOMER","COST_ANALYST","Cost analytics"),
("CUSTOMER","REPORTING_ADMIN","Reports and exports administration"),
("CUSTOMER","SECURITY_ADMIN","Company security"),
("CUSTOMER","AUDITOR","Company audit"),
("CUSTOMER","COMPLIANCE_OFFICER","Compliance monitoring"),
("CUSTOMER","PROCESS_HEAD","Process oversight"),
("CUSTOMER","PROCESS_ADMIN","Process administration"),
("CUSTOMER","PROCESS_TRANSPORT_ADMIN","Process transport operations"),
("CUSTOMER","SITE_ADMIN","Site administration"),
("CUSTOMER","SITE_TRANSPORT_ADMIN","Site transport"),
("CUSTOMER","SITE_SECURITY_ADMIN","Site security"),
("CUSTOMER","FACILITY_MANAGER","Facility/site coordination"),
("CUSTOMER","SITE_OPERATIONS_MANAGER","Site operations"),
("CUSTOMER","MANAGER","Team/process approvals"),
("CUSTOMER","TEAM_LEADER","Team transport"),
("CUSTOMER","SHIFT_SUPERVISOR","Shift operations"),
("CUSTOMER","TRAVEL_DESK_AGENT","Booking support"),
("CUSTOMER","TRANSPORT_HELPDESK_AGENT","Transport support"),
("CUSTOMER","BOOKING_COORDINATOR","Booking administration"),
("CUSTOMER","EXECUTIVE_ASSISTANT_BOOKER","Authorized executive booking"),
("CUSTOMER","EVACUATION_COORDINATOR","Emergency transport coordination"),
("CUSTOMER","EMPLOYEE","Own transport"),
("VENDOR","VENDOR_ADMIN","Vendor organization admin"),
("VENDOR","VENDOR_SUB_ADMIN","Delegated vendor administration"),
("VENDOR","VENDOR_OPERATIONS_MANAGER","Vendor operations"),
("VENDOR","VENDOR_DISPATCHER","Vendor dispatch"),
("VENDOR","VENDOR_FLEET_MANAGER","Vendor fleet"),
("VENDOR","VENDOR_DRIVER_MANAGER","Vendor drivers"),
("VENDOR","VENDOR_COMPLIANCE_MANAGER","Vendor compliance"),
("VENDOR","VENDOR_FINANCE","Vendor finance"),
("VENDOR","VENDOR_COORDINATOR","Trip coordination"),
("VENDOR","VENDOR_VIEWER","Vendor read-only"),
("EXTERNAL","DRIVER","Driver app"),
("EXTERNAL","DRIVER_SUPERVISOR","Driver supervision"),
("EXTERNAL","GUARD","Site verification"),
("EXTERNAL","GUARD_SUPERVISOR","Guard supervision"),
]
table(["Domain","Role","Responsibility"], roles)
page_break()

h1("7. Permission Catalogue")
permissions = [
("company.view","View companies"),("company.create","Create company"),("company.edit","Edit company"),
("company.activate","Activate company"),("company.suspend","Suspend company"),("company.archive","Archive company"),
("site.view","View sites"),("site.create","Create site"),("site.edit","Edit site"),
("process.view","View processes"),("process.create","Create process"),("process.edit","Edit process"),
("team.view","View teams"),("team.create","Create team"),("team.edit","Edit team"),
("user.view","View users"),("user.create","Create user"),("user.edit","Edit user"),("user.disable","Disable user"),
("user.change_role","Change role"),("user.assign_scope","Assign scope"),
("role.view","View roles"),("role.create","Create role"),("role.edit","Edit role"),("role.assign","Assign role"),
("permission.view","View permissions"),("permission.grant","Grant permission"),("permission.revoke","Revoke permission"),
("employee.view","View employee"),("employee.create","Create employee"),("employee.onboard","Onboard employee"),
("employee.edit","Edit employee"),("employee.offboard","Offboard employee"),("employee.restore","Restore employee"),
("employee.location.view","View employee location"),("employee.location.edit","Edit pickup/drop location"),
("booking.self","Book own transport"),("booking.team","Book team transport"),("booking.process","Book process transport"),
("booking.company","Book company-wide transport"),("booking.create","Create booking"),("booking.edit","Edit booking"),
("booking.cancel","Cancel booking"),("booking.bulk","Bulk booking"),
("trip.view","View trips"),("trip.create","Create trip"),("trip.assign","Assign trip"),("trip.unassign","Unassign driver"),
("trip.reassign","Change driver/whole-trip assignment"),("trip.vehicle_change","Change vehicle"),
("trip.passenger_add","Add passenger"),("trip.passenger_remove","Remove passenger"),
("trip.passenger_reassign","Move passenger(s)"),("trip.split","Split trip"),("trip.merge","Merge compatible trips"),
("trip.route_change","Change route"),("trip.schedule_change","Change trip schedule"),("trip.cancel","Cancel trip"),
("driver.view","View driver"),("driver.create","Create driver"),("driver.edit","Edit driver"),("driver.assign","Assign driver"),
("driver.activate","Manage driver availability"),("driver.suspend","Suspend driver"),("driver.document_manage","Manage driver docs"),
("vehicle.view","View vehicle"),("vehicle.create","Create vehicle"),("vehicle.edit","Edit vehicle"),("vehicle.assign","Assign vehicle"),
("vehicle.maintenance","Manage maintenance"),("vehicle.document_manage","Manage vehicle docs"),
("vehicle_qr.view","View vehicle QR"),("vehicle_qr.create","Generate vehicle-shift QR"),
("vehicle_qr.regenerate","Regenerate QR"),("vehicle_qr.revoke","Revoke QR"),
("vehicle_qr.delegate","Delegate QR generation"),("vehicle_qr.history","View QR history"),
("route.view","View routes"),("route.create","Create route"),("route.edit","Edit route"),("route.publish","Publish route"),
("schedule.view","View schedules"),("schedule.create","Create schedule"),("schedule.edit","Edit schedule"),
("roster.view","View roster"),("roster.manage","Manage roster"),
("noshow.view","View no-shows"),("noshow.mark","Mark no-show"),("noshow.approve","Approve no-show"),("noshow.override","Override no-show"),
("policy.view","View policy"),("policy.edit","Edit policy"),
("gps.view_live","View live GPS"),("gps.view_history","View GPS history"),("gps.view_driver","View driver GPS"),
("incident.view","View incidents"),("incident.create","Create incident"),("incident.manage","Manage incident"),
("safety.manage","Manage safety"),("emergency.manage","Manage emergency"),
("vendor.view","View vendor"),("vendor.create","Create vendor"),("vendor.edit","Edit vendor"),("vendor.compliance","Vendor compliance"),
("document.view","View documents"),("document.upload","Upload document"),("document.verify","Verify document"),
("document.reject","Reject/request re-upload"),("document.download","Download document"),("document.replace","Replace document"),
("document.archive","Archive document"),("document.restore","Restore document"),("document.history","View document history"),
("report.view","View reports"),("report.create","Create report"),("report.run","Run report"),("report.download","Download report"),
("report.export_excel","Export Excel"),("report.export_pdf","Export PDF"),("report.export_csv","Export CSV"),
("report.schedule","Schedule report"),("report.email","Email report"),("report.share","Share report"),
("report.manage_saved","Manage saved reports"),("report.manage_templates","Manage report templates"),
("report.cross_company","Cross-company reporting"),("report.cross_site","Cross-site reporting"),
("report.cross_process","Cross-process reporting"),("report.audit","Audit reporting"),("report.billing","Billing reporting"),
("report.data_quality","Data-quality reporting"),("report.admin","Report administration"),
("billing.internal.view","View internal SaaS billing"),("billing.internal.manage","Manage internal SaaS billing"),
("billing.external.view","View external transport billing"),("billing.external.ratecard","Manage transport rate cards"),
("billing.external.invoice","Manage external invoices"),("billing.external.approve","Approve external invoices"),
("billing.external.reconcile","Reconcile external billing"),("billing.export","Export finance"),
("audit.view","View audit"),("audit.export","Export audit"),("security.manage","Manage security"),
("session.revoke","Revoke session"),("mfa.manage","Manage MFA policy"),
("integration.view","View integrations"),("integration.manage","Manage integrations"),
("api_client.manage","Manage API clients"),("webhook.manage","Manage webhooks"),
("feature_flag.manage","Manage feature flags"),("system.health","View system health"),
("owner.manage","Manage owner identities"),
]
table(["Permission key","Human-readable responsibility"], permissions)
page_break()

h1("8. Super Admin Complete Access Matrix")
modules = [
("Platform","ALL","Create/manage platform objects and configuration"),
("Company Administration","ALL","Companies, sites, processes, teams and lifecycle"),
("Identity","ALL","Users, roles, permissions and scopes"),
("Transport","ALL","Employees, bookings, trips, schedules, rosters"),
("Dispatch","ALL","Assign, unassign, reassign, passenger movement, vehicle changes"),
("Fleet","ALL","Vendors, drivers, vehicles, assignments and maintenance"),
("Vehicle QR","ALL","Generate, regenerate, revoke, delegate, history"),
("GPS/Maps","ALL","Live/historical tracking, routes, geofences and map configuration"),
("Safety","ALL","SOS, incidents, female transport, emergency response"),
("Documents","ALL","Upload, verify, reject, replace, archive, download, history"),
("Reports","ALL","View, build, schedule, export, share and cross-tenant reporting"),
("Internal SaaS Billing","ALL","Plans, subscriptions, invoices, collections, revenue"),
("External Transport Billing","ALL","Rate cards, vendor invoices, reconciliation"),
("Integrations","ALL","HRMS, SSO, webhooks, APIs, notification providers"),
("Security","ALL","Sessions, MFA, security events and access tools"),
("Audit","ALL","All audit trails and exports"),
("System","ALL","Health, jobs, queues, feature flags, configuration"),
]
table(["Module","SUPER_ADMIN","Capability"], modules)
para("The same all-functional access applies to MOVEINSYNC_OWNER and SAAS_OWNER, with owner hierarchy controls determining who may manage owner-level identities.")
page_break()

h1("9. Company and Tenant Management")
h2("9.1 Company fields")
table(["Field","Required","Notes"], [
("Legal Name","Yes","Unique legal entity record"),
("Display Name","Yes","UI name"),
("Company Code","Yes","Immutable business identifier after activation unless controlled migration"),
("Billing Model","Yes","COMPANY / SITE / PROCESS"),
("Primary Contact","Yes","Customer owner/contact"),
("Office Email","Yes","Verified"),
("Mobile","Yes","Verified"),
("Alternate Mobile","No","Validated"),
("Address","Yes","Registered/operational address"),
("Country","Yes","Country/region"),
("Timezone","Yes","Default operational timezone"),
("Contract Start","No","Commercial metadata"),
("Contract End","No","Commercial metadata"),
("Default Transport Policy","Yes","Initial policy"),
("Default Safety Policy","Yes","Initial safety baseline"),
("Status","Yes","ACTIVE / SUSPENDED / ARCHIVED"),
])
h2("9.2 Company creation transaction")
code("""validate input
validate unique legal name / company code
create company
create default transport policy
create default safety policy
create billing configuration
create initial site/process scaffolding if selected
optionally invite first admin
write audit
commit
""")
para("Failure must not leave a half-created tenant. Invitation failure should not roll back the already valid customer record; show account status as invitation pending and allow resend.")
h2("9.3 Company lifecycle")
bl([
    "Draft - being configured before activation.",
    "Active - normal operations.",
    "Suspended - operational access blocked according to policy while data is retained.",
    "Archived - no new operational transactions; historical access subject to policy.",
    "Restored - controlled return from archive/suspension where authorized."
])
page_break()

h1("10. Site, Process, Team and Shift Hierarchy")
para("Company Code, Site Code and Process Code are first-class tracking dimensions. Do not store them as comma-separated text. Use normalized foreign-key relationships and assignment tables.")
table(["Object","Example code","Important functions"], [
("Company","CMP-00127","Tenant identity and billing relationship"),
("Site","SITE-BLR01","Operational location and site policy"),
("Process","PROC-P01","Business process / cost scope"),
("Team","TEAM-CC01","Manager/team grouping"),
("Shift","SHIFT-1830","Recurring operational time window"),
])
h2("10.1 Multi-scope users")
para("Directors, Transport Admins and other authorized roles can have multiple site and process scopes. Scope assignment is many-to-many and independently auditable.")
h2("10.2 Overnight shifts")
para("A shift such as 18:30-03:30 is one logical overnight shift. Store local-date semantics explicitly so trips do not break at midnight.")
page_break()

h1("11. User and Authentication Management")
h2("11.1 Supabase Auth")
para("Use Supabase Auth for authentication, sessions, password recovery and email verification. The application database remains the source of truth for business roles, scopes and permissions.")
h2("11.2 Login flows")
bl([
    "Invitation/first-time activation: invite → verify email/OTP as configured → establish password/session → load effective access.",
    "Normal login: email/password or approved authentication method → session → effective access.",
    "Password reset: recovery request → verification → new password → session security handling.",
    "MFA: required for owner-level and privileged roles according to policy."
])
h2("11.3 Separation of OTPs")
table(["Mechanism","Purpose","Authority"], [
("Login/recovery OTP","Authentication/account recovery","Supabase Auth"),
("Trip/boarding OTP","Operational trip verification","Move In Sync transport service"),
("QR","Vehicle/duty identity and operational verification","Transport operations"),
])
page_break()

h1("12. Employee Transport Master")
para("The employee transport record is independent from HRMS. HRMS can be an integration source, but it cannot silently control transport state.")
table(["Field","Required"], [
("Full Name","Yes"),("Employee ID","Yes"),("Company Code / Company ID","Yes"),("Site Code / Site ID","Yes"),
("Process Code / Process ID","Yes"),("Office Email","Yes"),("Mobile Number","Yes"),("Alternate Mobile","No"),
("Gender where safety policy requires","Conditional"),("Designation","No"),("Department","No"),("Manager ID","No"),
("Team ID","No"),("Shift / Shift Group","No"),("Pickup Latitude","Yes"),("Pickup Longitude","Yes"),
("Pickup Address","Yes"),("Office/Site Latitude","Yes"),("Office/Site Longitude","Yes"),("Office/Site Address","Yes"),
("Transport Eligible","Yes"),("Transport Status","Yes"),("Effective From","Yes"),("Effective To","No"),
("Emergency Contact","Policy dependent"),("Accessibility requirements","Policy dependent")
])
h2("12.1 Onboard")
bl(["Transport Admin opens Employees → Add Employee.", "System restricts company/site/process selectors to administrator scope.", "Validate identity and contact information.", "Validate pickup coordinates/address.", "Preview record.", "Confirm ONBOARD.", "Persist and audit.", "Optionally initiate application invitation."])
h2("12.2 Offboard")
bl(["Require effective date/time and reason.", "Show future bookings and trips.", "Require explicit disposition for affected future transport.", "Set transport status OFFBOARD.", "Block new bookings after effective time.", "Preserve historical trips, invoices, no-shows and audit.", "Write immutable audit event."])
page_break()

h1("13. Pickup Locations and Geodata")
h2("13.1 Employee pickup points")
para("An employee can have multiple approved pickup points.")
table(["Field","Purpose"], [
("Pickup Code","Human-readable location identifier"),
("Label","Home / Main Gate / Temporary"),
("Address","Human-readable location"),
("Latitude/Longitude","Routing point"),
("Geofence radius","Operational tolerance"),
("Landmark","Driver-friendly guidance"),
("Instructions","Pickup notes"),
("Effective from/to","Validity"),
("Verified","Address/data-quality status"),
("Primary","Default selection"),
])
h2("13.2 Map picker")
para("Provide a map-based point picker, address search, reverse geocoding and validation. Users should not need to manually type coordinates.")
page_break()

h1("14. Fleet and Vendor Master")
h2("14.1 Vendor")
bl(["Vendor legal/display name.", "Vendor Code.", "Contacts.", "Contract dates.", "Assigned companies/sites/processes.", "Rate cards.", "SLA.", "Compliance documents.", "Bank/payment information according to secure finance access.", "Status."])
h2("14.2 Driver")
bl(["Driver Code.", "Vendor.", "Identity/contact.", "Licence and permits.", "Background verification.", "Training.", "Assigned operating areas.", "Current state.", "Documents.", "Mobile/device registration."])
h2("14.3 Vehicle")
bl(["Vehicle Code / registration number.", "Vehicle type.", "Seated/accessible capacity.", "Vendor.", "Make/model/year.", "Fuel/EV type.", "GPS device.", "Registration/fitness/insurance/permit.", "Maintenance state.", "Emergency equipment.", "Status."])
page_break()

h1("15. Document Management - Complete Functional Specification")
para("Documents are a reusable platform capability. Every document-enabled entity must visibly expose a Documents section where the user's permissions allow it.")
h2("15.1 Where Upload appears")
table(["Entity","Example document","UI location"], [
("Driver","Driving licence","Driver → Documents"),
("Vehicle","Insurance","Vehicle → Documents"),
("Vendor","Contract","Vendor → Documents"),
("Employee","Approved identity/supporting document where policy requires","Employee → Documents"),
("Incident","Photo/evidence","Incident → Evidence"),
("Invoice","Supporting bill","Invoice → Attachments"),
("Company","Contract/legal document","Company → Documents"),
("Site","Permit/site document","Site → Documents"),
])
h2("15.2 Document actions")
bl(["Upload Document", "Preview", "Download", "Verify", "Reject", "Request Re-upload", "Replace / New Version", "Archive", "Restore", "View History"])
h2("15.3 Document lifecycle")
code("""DRAFT
  |
  v
UPLOADED
  |
  v
PENDING_REVIEW
  +--> VERIFIED
  +--> REJECTED --> REUPLOAD
  |
  +--> EXPIRED
  |
  +--> REPLACED / ARCHIVED
""")
h2("15.4 Document record")
table(["Field","Description"], [
("Document Code","Human-readable stable identifier"),
("Entity Type/ID","Parent entity"),
("Document Type","Licence, insurance, contract, etc."),
("Version","Monotonic version"),
("File reference","Private storage object reference"),
("Issue date","Document issue date"),
("Expiry date","Optional/conditional"),
("Status","Lifecycle state"),
("Verification status","Pending/verified/rejected"),
("Verified by/at","Review evidence"),
("Rejection reason","Required when rejected"),
("Uploaded by/at","Traceability"),
])
h2("15.5 Storage")
para("Use private Supabase Storage buckets. Do not expose unrestricted public URLs for sensitive documents. Downloads require authorization and a short-lived controlled access mechanism.")
page_break()

h1("16. QR and Vehicle Duty System")
h2("16.1 Authority")
table(["Role","Generate QR","Regenerate/Revoke","Use/scan"], [
("SUPER_ADMIN / OWNER","Yes","Yes","Yes"),
("TRANSPORT_ADMIN","Yes","Yes","Yes"),
("TRANSPORT_COORDINATOR","Only if delegated","Only if delegated","Yes"),
("DRIVER","No","No","Operational use only"),
("MANAGER","No","No","No"),
("FINANCE","No","No","No"),
("EMPLOYEE","No","No","No"),
])
h2("16.2 QR belongs to a duty")
code("""Vehicle
  -> Vehicle Duty
      -> Company
      -> Site
      -> Process
      -> Shift
      -> Service date
      -> Driver assignment
      -> QR instance
""")
h2("16.3 QR record")
table(["Field","Description"], [
("QR Code","Human-readable identifier"),
("Token hash","Protected QR secret representation"),
("Vehicle Duty","Context binding"),
("Driver assignment","Operational identity"),
("Valid from/to","Time window"),
("Status","ACTIVE / REVOKED / EXPIRED"),
("Generated by","Actor"),
("Generated at","Timestamp"),
("Revoked by/at","Audit"),
])
h2("16.4 QR events")
bl(["Created", "Regenerated", "Revoked", "Scanned", "Invalid", "Expired", "Replay detected", "Wrong vehicle", "Wrong duty", "Wrong scope"])
page_break()

h1("17. Booking Engine")
h2("17.1 Booking types")
bl(["Self", "Team", "Process", "Company-wide", "Recurring", "Bulk", "Standing", "Emergency", "Guest/visitor where configured"])
h2("17.2 Booking validation")
bl(["Employee active and transport eligible.", "Site/process active.", "Booking window open.", "Cutoff not exceeded.", "Shift exists.", "Pickup location approved.", "Capacity/policy permits.", "Requester has permission and scope.", "Employee not offboarded at effective booking time.", "Duplicate booking rules checked."])
h2("17.3 Booking codes")
para("Every booking receives a stable Booking Code used by employees, operations and support. The database UUID remains the technical key.")
page_break()

h1("18. Trip Engine")
h2("18.1 Trip identity")
para("Every trip must receive a human-readable Trip Code such as TRP-20260905-001274. The Trip Code is searchable, reportable, visible to authorized users and used in operational communication.")
h2("18.2 Trip lifecycle")
code("""CREATED
  -> ASSIGNMENT_PENDING
  -> ASSIGNED
  -> DISPATCHED
  -> DRIVER_EN_ROUTE
  -> DRIVER_ARRIVED
  -> BOARDING
  -> STARTED
  -> IN_TRANSIT
  -> DROP_COMPLETED
  -> COMPLETED

Exceptions:
CANCELLED
NO_SHOW
REASSIGNED
ABORTED
EMERGENCY
GPS_EXCEPTION
""")
h2("18.3 Trip code is not enough")
para("Trip Code identifies the trip. Passenger ID identifies an individual passenger record. Assignment ID identifies a specific driver/vehicle assignment. Duty Code identifies the operational vehicle duty. Keep all levels separate.")
page_break()

h1("19. Passenger-Level Trip Operations")
para("This is a mandatory real-world capability. A trip is a container of passengers, not an indivisible record.")
h2("19.1 Example")
code("""TRP-20260905-001274
Driver: Ravi
Vehicle: KA01AB1234
Passengers:
  EMP-001
  EMP-002
  EMP-003
  EMP-004
""")
para("Transport Admin can select EMP-004 and choose Move Passenger. The system finds eligible destination trips and shows capacity, route/time impact, safety impact and ETA impact before applying the move.")
h2("19.2 Passenger actions")
bl(["Remove passenger.", "Move passenger to another trip.", "Move multiple selected passengers.", "Add passenger to existing trip.", "Change one passenger pickup.", "Cancel one passenger while leaving the trip active.", "Mark no-show.", "Restore passenger where policy allows."])
h2("19.3 Source/destination behavior")
para("Moving a passenger creates a source removal event and destination addition event. It must not alter unrelated passengers. The source trip can continue if it remains operationally valid. The destination trip must be revalidated before commit.")
page_break()

h1("20. Driver and Vehicle Reassignment")
h2("20.1 Transport Admin controls")
bl(["Unassign current driver.", "Assign a different eligible driver.", "Replace vehicle.", "Replace driver and vehicle together.", "Transfer duty where policy permits.", "Recalculate route/ETA/capacity.", "Recheck documents and driver state.", "Notify affected parties."])
h2("20.2 State-aware editing")
table(["Trip state","Driver change","Vehicle change","Passenger move"], [
("CREATED","Allowed","Allowed","Allowed"),
("ASSIGNED","Allowed","Allowed","Allowed"),
("DISPATCHED","Allowed with impact/revalidation","Allowed with impact/revalidation","Allowed with strong checks"),
("DRIVER_ARRIVED","Controlled/exception workflow","Controlled/exception workflow","Controlled/exception workflow"),
("STARTED","Normally blocked; emergency override only","Normally blocked; emergency override only","Normally blocked; emergency override only"),
("COMPLETED","No","No","No; history only"),
])
h2("20.3 Assignment history")
para("Never overwrite historical driver/vehicle assignment records. Create assignment versions and events.")
page_break()

h1("21. Split and Merge Trips")
h2("21.1 Split")
para("Split is used when one trip must become multiple trips because of capacity, route, time, safety or operational constraints.")
bl(["Select passengers.", "Choose split reason.", "Show resulting trip count.", "Select new vehicle/driver if required.", "Preview route/time/capacity.", "Confirm.", "Generate new Trip Code(s).", "Preserve source Trip Code history."])
h2("21.2 Merge")
para("Merge is allowed only for compatible trips. Compatibility checks include company/site/process policy, pickup windows, route geography, capacity, vehicle type, female safety rules, driver hours and operational state.")
page_break()

h1("22. Dispatch Control Room")
h2("22.1 Primary layout")
code("""+-------------------------------------------------------------+
| LIVE OPERATIONS                                             |
| Unassigned | Delayed | SOS | GPS stale | No-show | Capacity |
+---------------------------+---------------------------------+
|                           | Exceptions                      |
|         LIVE MAP          | - driver unavailable            |
|                           | - wrong route                   |
|                           | - document expired              |
|                           | - safety                         |
+---------------------------+---------------------------------+
| Trip Queue | Driver Availability | Vehicle Availability      |
+-------------------------------------------------------------+
""")
h2("22.2 Operator actions")
bl(["Open trip.", "Assign driver.", "Change vehicle.", "Move passenger.", "Open driver.", "Open vehicle.", "View route.", "View GPS.", "Resolve exception.", "View audit/timeline."])
h2("22.3 Map/list synchronization")
para("Selecting a trip highlights the vehicle and route. Selecting a vehicle highlights the current trip. Selecting an exception opens the relevant object without losing filter context.")
page_break()

h1("23. Route and Schedule Management")
h2("23.1 Route design")
bl(["Route Code.", "Route name.", "Origin and destination.", "Stops.", "Stop order.", "Dwell time.", "Geofence.", "Pickup window.", "Capacity.", "Route version.", "Published state."])
h2("23.2 Schedule")
bl(["One-time.", "Recurring.", "Shift-linked.", "Service windows.", "Booking cutoff.", "Dispatch cutoff.", "Holiday exceptions.", "Blackout dates.", "Driver break windows."])
h2("23.3 Versioning")
para("Published routes/schedules are versioned. A new version should not rewrite the historical route used by completed trips.")
page_break()

h1("24. Roster and Shift Planning")
bl(["Shift templates.", "Employee-shift assignments.", "Roster calendar.", "Site/process shift coverage.", "Holiday calendar.", "Overtime transport.", "Shift swap requests.", "Roster import.", "Conflicting shift detection.", "Midnight-safe logic."])
para("Roster planning should show expected transport demand so planners can identify under-capacity shifts before publishing.")
page_break()

h1("25. Driver State and Fatigue")
table(["State","Dispatch eligibility"], [
("OFFLINE","No"),
("AVAILABLE","Yes"),
("ON_TRIP","No"),
("BREAK","No"),
("SLEEPY","No"),
("EMERGENCY","No"),
("SUSPENDED","No"),
])
bl(["Driver can mark break when no trip is active.", "Driver can mark sleepy before a trip.", "Sleepy state removes dispatch eligibility.", "Emergency state triggers incident handling.", "Driver preferred areas influence ranking but do not override safety."])
page_break()

h1("26. Driver App")
h2("26.1 Main screen")
bl(["Current duty.", "Vehicle.", "Current Trip Code.", "Next pickup.", "Passenger count.", "GPS status.", "Network status.", "Primary trip action."])
h2("26.2 Actions")
bl(["Go Active.", "Break.", "Sleepy.", "Start Trip.", "Arrived.", "No-show.", "Complete Trip.", "Emergency/SOS.", "Trip history."])
h2("26.3 Mobile principles")
bl(["Large touch targets.", "Minimal typing.", "Offline-safe queue for eligible actions.", "Clear status labels.", "Persistent emergency access.", "No administrative clutter."])
page_break()

h1("27. Employee App")
bl(["My upcoming trip.", "Book transport.", "Cancel permitted booking.", "Pickup location.", "Vehicle/driver details where policy permits.", "ETA.", "Trip history.", "Approved pickup points.", "Report issue.", "Safety/emergency."])
para("Employee should never have to understand Company Code, Site Code or Process Code to complete normal self-service actions. Those values can remain available as secondary context.")
page_break()

h1("28. Manager and Team Leader Experience")
bl(["Team transport dashboard.", "Upcoming team trips.", "Team bookings.", "No-shows.", "Approval queue where configured.", "Transport issue reporting.", "Team-level reports where permitted."])
para("Manager and Team Leader do not receive vehicle QR generation merely because they can view team transport.")
page_break()

h1("29. Guard and Site Security")
bl(["Vehicle verification.", "Driver verification.", "Trip verification.", "Expected passenger count.", "Site entry/exit.", "QR/identity verification.", "Incident report.", "Security notes."])
para("Guard sees only operational verification data necessary for safe site access. Do not expose finance, unrelated employee records or global transport configuration.")
page_break()

h1("30. Safety and Emergency Management")
bl(["SOS.", "Emergency trip flag.", "Female transport policies.", "Safe pickup rules.", "Escort requirement.", "Restricted driver rules.", "Route deviation alert.", "Missed check-in.", "Emergency contacts.", "Incident case management.", "Emergency response timeline."])
para("Safety rules have precedence over optimization. A cheaper or faster route is never chosen if it violates an active safety rule.")
page_break()

h1("31. Incident Management")
h2("31.1 Incident lifecycle")
code("""OPEN
 -> TRIAGED
 -> ASSIGNED
 -> INVESTIGATING
 -> ACTION_REQUIRED
 -> RESOLVED
 -> CLOSED
""")
h2("31.2 Incident record")
bl(["Incident Code.", "Severity.", "Category.", "Time.", "Company/site/process.", "Trip/vehicle/driver/passenger references.", "Description.", "Evidence/documents.", "Actions.", "Owner.", "Resolution.", "Closure time."])
page_break()

h1("32. GPS Platform")
h2("32.1 GPS event")
table(["Field","Purpose"], [
("Driver ID","Identity"),
("Device ID","Source"),
("Trip ID","Current trip"),
("Timestamp","Event time"),
("Latitude / Longitude","Position"),
("Accuracy","Quality"),
("Speed / heading","Optional motion data"),
("Connectivity","Network state"),
("Sequence number","Dedupe/order"),
])
h2("32.2 GPS states")
bl(["LIVE.", "STALE.", "OFFLINE.", "INVALID.", "LOW ACCURACY."])
h2("32.3 Ingestion")
para("Use authenticated ingestion. Accept retries safely with a unique device/sequence or equivalent idempotency key. Mobile queues events during intermittent connectivity and retries safely.")
page_break()

h1("33. Map Platform and API Keys")
para("Use a MapProvider abstraction so the transport domain is not locked to one map vendor.")
h2("33.1 Primary provider")
para("Google Maps Platform is the primary production provider for web maps, geocoding, address validation, routing and route matrices. Google's current security guidance recommends application/API restrictions, separate keys by application, client/server separation and disabling unused services. citeturn749546search0")
h2("33.2 Key model")
table(["Credential","Use","Restriction"], [
("NEXT_PUBLIC_GOOGLE_MAPS_BROWSER_KEY","Web Maps JavaScript and allowed browser map services","Website/referrer restriction + API restriction"),
("GOOGLE_MAPS_GEOCODING_KEY","Server geocoding/reverse geocoding","Server/IP restriction where applicable + API restriction"),
("GOOGLE_MAPS_ADDRESS_VALIDATION_KEY","Server address validation","Server/IP restriction where applicable + API restriction"),
("GOOGLE_MAPS_ROUTES_KEY","Server routing and travel-time matrix","Server/IP restriction where applicable + API restriction"),
("GOOGLE_MAPS_ROADS_KEY","Optional GPS road matching if required","Server/IP restriction + API restriction"),
("GOOGLE_MAPS_STATIC_KEY","Static maps for reports/PDF where required","Application/API restriction"),
("GOOGLE_ROUTE_OPTIMIZATION_SERVICE_ACCOUNT","Fleet optimization","Server-side identity/credential only"),
])
h2("33.3 Current Google capabilities used by the design")
para("Google's current Routes API documents Compute Route Matrix for up to 625 route elements. The Address Validation API validates address components and returns geocoding/deliverability information. citeturn433501search2turn433501search1")
h2("33.4 Secret handling")
bl(["Never put server keys in source code.", "Never paste production secrets into Mimo prompts.", "Use environment/secret storage.", "Use different credentials/projects for dev, staging and production where practical.", "Audit key usage and rotate with migration planning.", "Proxy server-side web-service calls rather than allowing arbitrary client requests."])
page_break()

h1("34. Map Features")
bl(["Map rendering.", "Address autocomplete/search.", "Geocoding.", "Reverse geocoding.", "Address validation.", "Route calculation.", "Route matrix.", "ETA.", "Geofences.", "GPS replay.", "Live vehicle map.", "Route deviation.", "Pickup-point map picker.", "Route designer.", "Vehicle proximity search.", "Optimization-ready provider adapter."])
page_break()

h1("35. Dispatch Optimization")
h2("35.1 Inputs")
bl(["Driver state.", "Vehicle capacity.", "Vehicle type.", "Current position.", "Shift window.", "Driver hours.", "Passenger pickup windows.", "Route compatibility.", "Vendor rules.", "Safety rules.", "Document compliance.", "Cost/rate card.", "Preferred operating area."])
h2("35.2 Ranking")
para("Optimization produces ranked suggestions. It does not silently change a trip without the workflow permitted by policy.")
h2("35.3 Safety precedence")
code("""Safety constraints
   > compliance constraints
   > legal/working-hour constraints
   > capacity constraints
   > service-level constraints
   > cost / distance optimization
""")
page_break()

h1("36. Reports and Reporting Platform")
para("Reporting is a product area, not a single button.")
h2("36.1 Report permissions")
bl(["report.view", "report.create", "report.run", "report.download", "report.export_excel", "report.export_pdf", "report.export_csv", "report.schedule", "report.email", "report.share", "report.manage_saved", "report.manage_templates", "report.cross_company", "report.cross_site", "report.cross_process", "report.audit", "report.billing", "report.data_quality", "report.admin"])
h2("36.2 Report workflow")
code("""Choose report
  -> choose filters
  -> choose authorized fields
  -> preview
  -> run
  -> export format
  -> background job for large output
  -> secure download
  -> audit
""")
page_break()

h1("37. Report Catalogue")
report_names = [
"Employee Master","Employee Transport Eligibility","Booking Detail","Trip Detail","Trip Passenger Movement",
"Dispatch Performance","Driver Availability","Driver Productivity","Vehicle Utilization","Fleet Maintenance",
"Route Utilization","Schedule Adherence","No-Show","GPS Health","ETA and Delay","Safety and Incident",
"Female Transport Safety","Vendor SLA","Vendor Cost","Cost Per Trip","Cost Per Employee","Rate Card",
"External Invoice","Billing Reconciliation","Internal SaaS Billing","Collections","Revenue","Document Compliance",
"Document Expiry","Audit","Access and Permission","Security Event","Integration/Webhook","Feature Flag Change",
"System Health","Data Quality","Custom Report"
]
for x in report_names:
    bullet(x)
page_break()

h1("38. Report Builder")
bl(["Choose a subject area.", "Select business fields.", "Apply scope-safe filters.", "Group and sort.", "Preview.", "Save.", "Export.", "Schedule if permitted."])
para("Do not expose arbitrary SQL to normal users. The report builder works from a controlled semantic model and approved fields.")
page_break()

h1("39. Report Downloads and Scheduling")
h2("39.1 Download formats")
bl(["Excel (.xlsx).", "PDF.", "CSV."])
h2("39.2 Large report jobs")
bl(["Create export job.", "Show job status.", "Keep browser responsive.", "Securely store output.", "Expire download link after configured period.", "Audit requester, report, format, filters and row count.", "Retain failure reason."])
h2("39.3 Scheduled reports")
bl(["Daily/weekly/monthly as configured.", "Authorized recipients only.", "Stored filters and columns.", "Automatically stop if access is revoked.", "Run history and failure alerts."])
page_break()

h1("40. Billing Architecture")
para("Keep internal Move In Sync SaaS billing and external customer/vendor transport billing in separate domains.")
h2("40.1 Internal SaaS billing")
bl(["Customer subscription.", "Plan.", "Modules.", "Usage/seats where applicable.", "Contracted price.", "SaaS invoice.", "Tax.", "Credits/debits.", "Collections.", "Payment status.", "Revenue."])
h2("40.2 External transport billing")
bl(["Vendor contract.", "Rate card.", "Trip charge.", "Per-km/per-hour/per-trip.", "Waiting.", "Tolls/parking.", "Special services.", "Vendor invoice.", "Validation.", "Dispute.", "Credit/debit.", "Reconciliation.", "Payment state."])
page_break()

h1("41. Finance Roles and Permissions")
table(["Role","Primary scope","Can approve payment?","Report download"], [
("PLATFORM_FINANCE_ADMIN","SaaS/platform","Per governance policy","All authorized platform finance"),
("PLATFORM_FINANCE_VIEWER","SaaS/platform","No","Yes, read-only scope"),
("PLATFORM_BILLING_ADMIN","Customer subscriptions","Per policy","All subscription reports"),
("FINANCE_ADMIN","Customer transport","No unless separately granted","All company finance"),
("FINANCE_APPROVER","Customer transport","Yes where assigned","Authorized finance"),
("FINANCE_VIEWER","Customer transport","No","Authorized finance"),
("COST_ANALYST","Cost analytics","No","Cost reports"),
("VENDOR_FINANCE","Vendor scope","Per vendor process","Vendor finance"),
])
para("SUPER_ADMIN, MOVEINSYNC_OWNER and SAAS_OWNER have complete functional access across both billing domains, subject to secure approval workflow and audit.")
page_break()

h1("42. Rate Cards")
bl(["Versioned rate cards.", "Effective from/to.", "Per trip.", "Per km.", "Per hour.", "Per vehicle/day.", "Per passenger.", "Night surcharge.", "Holiday surcharge.", "Waiting charge.", "Toll/parking.", "Emergency/special vehicle.", "Vendor-specific overrides."])
para("A rate card used by a completed trip must remain historically reconstructable.")
page_break()

h1("43. Invoice and Reconciliation")
h2("43.1 Invoice lifecycle")
code("""DRAFT -> GENERATED -> SUBMITTED -> UNDER_REVIEW
      -> APPROVED -> DISPUTED / REJECTED
      -> POSTED -> PAID / PARTIALLY_PAID
""")
h2("43.2 Traceability")
para("Every invoice line must trace to underlying trips, services, rate-card versions and approved adjustments.")
h2("43.3 Reconciliation")
bl(["Expected operational charges.", "Vendor invoice lines.", "Variance.", "Missing trip.", "Duplicate line.", "Rate mismatch.", "Tax mismatch.", "Manual adjustment with reason."])
page_break()

h1("44. Notifications and Communication")
bl(["In-app.", "Email.", "Push.", "SMS.", "WhatsApp through provider adapter.", "Webhook."])
h2("44.1 Notification events")
bl(["Trip assigned.", "Driver changed.", "Vehicle changed.", "Passenger moved.", "Trip delayed.", "Driver arrived.", "Booking cancelled.", "Document expiring.", "Document rejected.", "Driver unavailable.", "Emergency.", "Invoice generated.", "Report ready."])
page_break()

h1("45. HRMS and Enterprise Integrations")
bl(["HRMS employee feed.", "Roster feed.", "SSO/identity.", "ERP/finance.", "Vendor APIs.", "GPS providers.", "Email.", "SMS.", "Push.", "WhatsApp.", "BI/data warehouse.", "Webhooks."])
para("Integrations use adapters. An integration failure must be visible and retryable; it must not silently mutate transport state.")
page_break()

h1("46. Webhooks and API Clients")
h2("46.1 API clients")
bl(["Client code.", "Name.", "Purpose.", "Scopes.", "Status.", "Created by.", "Last used.", "Credential rotation.", "Audit."])
h2("46.2 Webhooks")
bl(["Endpoint.", "Events.", "Signing secret.", "Status.", "Retry policy.", "Dead-letter/error state.", "Delivery history.", "Replay controls where safe."])
page_break()

h1("47. Policy Engine")
para("Policies are configuration data, not hard-coded branches scattered across the UI.")
bl(["Booking cutoff.", "No-show grace.", "QR validity/rotation.", "Female transport rules.", "Driver sleepy policy.", "Maximum trip duration.", "Pickup deviation tolerance.", "Location-change approval.", "Document expiry blocking.", "Report retention.", "Notification channels."])
page_break()

h1("48. Feature Flags")
bl(["Platform-wide.", "Company.", "Site.", "Process.", "Role.", "Percentage/canary.", "Environment."])
para("Feature flags control rollout, not authorization. Authorization remains governed by permissions. A disabled feature must fail predictably and explain its state where appropriate.")
page_break()

h1("49. Access Simulator")
para("Owner-level/security users can simulate access without mutating production data.")
table(["Input","Example"], [
("User","USR-001"),
("Role","TRANSPORT_ADMIN"),
("Company","CMP-001"),
("Site","SITE-BLR01"),
("Process","PROC-P01"),
("Team","TEAM-01"),
("Action","trip.passenger_reassign"),
("Resource","TRP-20260905-001274"),
])
para("Result must state ALLOWED or DENIED and explain the exact reason, e.g. OUT_OF_SCOPE, PERMISSION_OFF, ROLE_NOT_ALLOWED, RESOURCE_STATE_BLOCKED, POLICY_BLOCKED.")
page_break()

h1("50. Security Centre")
bl(["Active sessions.", "Login history.", "Failed logins.", "MFA status.", "Privileged changes.", "API credentials.", "Security events.", "Suspicious QR use.", "Document access.", "Export access.", "Session revocation."])
page_break()

h1("51. Audit and Evidence")
para("Sensitive mutations are immutable evidence events.")
table(["Audit field","Requirement"], [
("Actor","User/service identity"),
("Action","Stable action key"),
("Object","Entity and ID"),
("Timestamp","Server timestamp"),
("Before","Previous state when relevant"),
("After","New state when relevant"),
("Reason","Required for configured sensitive actions"),
("Scope","Resolved company/site/process"),
("Request ID","Correlation"),
("IP/device","Security metadata when appropriate"),
])
h2("51.1 Human readable audit")
para("Provide a readable timeline as well as technical detail. Example: “02:14 AM - Transport Admin moved EMP-004 from TRP-001274 to TRP-001301 because the destination trip had a compatible route and available capacity.”")
page_break()

h1("52. Universal Search")
bl(["Trip Code.", "Booking Code.", "Employee ID/name.", "Vehicle registration/code.", "Driver code/name.", "Vendor code.", "Route Code.", "Incident Code.", "Invoice Code.", "Document Code."])
para("Search is scope-aware and returns grouped results with status and context.")
page_break()

h1("53. User-Friendly UX Master Rules")
h2("53.1 No artificial page limit")
para("There is no page-count limit. Build however many screens, flows and detail views are necessary.")
h2("53.2 Complexity should be organized, not deleted")
bl(["Use progressive disclosure.", "Use drawers for focused context.", "Use wizards for long forms.", "Use tabs for related information.", "Use clear breadcrumbs.", "Use contextual actions.", "Use human-readable names."])
h2("53.3 Technical keys stay behind the UI")
table(["Technical","User-facing"], [
("trip.passenger_reassign","Move passenger"),
("vehicle_qr.create","Generate vehicle QR"),
("report.export_excel","Download Excel"),
("employee.offboard","Offboard employee"),
("billing.external.reconcile","Reconcile transport billing"),
])
page_break()

h1("54. Universal Detail Page Pattern")
code("""Header
  Name / Code / Status / Primary action

Summary
  Key facts and current state

Context
  Company / Site / Process / Team / Shift

Actions
  Only actions allowed for current user + object state

Tabs
  Overview | Activity | Documents | History | Related

Timeline
  Business events

Audit
  Who / when / before / after / reason
""")
para("Users should not have to open five unrelated pages for a single operational decision.")
page_break()

h1("55. Smart Forms")
bl(["Required fields clearly marked.", "Plain-language explanations.", "Searchable selectors.", "Map picker.", "Address validation.", "Timezone-aware date/time.", "Draft/save where appropriate.", "Preserve entered values after API failure.", "Client and server validation."])
page_break()

h1("56. Empty, Loading, Error and Forbidden States")
table(["State","Required behavior"], [
("Loading","Skeleton/spinner with preserved context"),
("Empty","Explain what is missing and present next safe action"),
("Error","Explain what failed and offer retry"),
("403","Explain that access is restricted without leaking data"),
("Offline","Show connection state and safe offline behavior"),
("Processing","Show background job state"),
("Success","Confirm the actual saved operation"),
("Partial success","Show successes and failures distinctly"),
])
page_break()

h1("57. Human-Friendly Error Messages")
bl([
    "“Driver cannot be assigned because the required licence expired on 02 Sep 2026.”",
    "“This process is outside your assigned process scope.”",
    "“Trip is already started; passenger movement requires an emergency override workflow.”",
    "“Document uploaded successfully and is pending verification.”",
    "“Report contains 42,817 rows and will be generated in the background.”",
    "“GPS has not been received for 8 minutes; vehicle is marked STALE.”",
])
page_break()

h1("58. Mobile UX")
h2("58.1 Driver")
bl(["Duty-first.", "Large controls.", "Very few text inputs.", "Offline queue.", "GPS/network visibility.", "Emergency action."])
h2("58.2 Employee")
bl(["Book.", "Upcoming.", "Pickup.", "Trip status.", "History.", "Help.", "Safety."])
h2("58.3 Guard")
bl(["Scan/verify.", "Expected trip.", "Allow/deny entry.", "Incident."])
page_break()

h1("59. Accessibility")
bl(["Keyboard navigation.", "Visible focus.", "Meaning not communicated by color alone.", "Readable typography.", "Accessible form errors.", "Labels on every control.", "Screen-reader-friendly status messages.", "Touch targets appropriate for mobile.", "Avoid unnecessary motion."])
page_break()

h1("60. Data Model - Core Tables")
tables = [
"users","user_profiles","roles","permissions","role_permissions","user_roles","user_permission_overrides",
"owner_identities","organizations","companies","company_sites","processes","teams","team_members",
"user_company_scopes","user_site_scopes","user_process_scopes","user_team_scopes",
"employees","employee_transport_profiles","employee_locations","employee_pickup_points",
"shift_templates","shifts","employee_shift_assignments",
"vendors","vendor_company_scopes","vendor_site_scopes","vendor_process_scopes","vendor_documents",
"drivers","driver_documents","driver_assignments","driver_status_events","driver_preferred_areas",
"vehicles","vehicle_documents","vehicle_assignments","vehicle_duties","vehicle_qr_instances","vehicle_qr_events",
"routes","route_stops","route_versions","geofences","schedules",
"bookings","booking_passengers","booking_events",
"trips","trip_passengers","trip_assignments","trip_events",
"gps_events","gps_latest_state","no_show_events",
"incidents","incident_events",
"documents","document_versions","document_reviews",
"transport_policies","safety_policies","policy_versions",
"rate_cards","rate_card_versions","invoices","invoice_lines","reconciliations",
"notifications","notification_templates","notification_deliveries",
"api_clients","webhooks","webhook_deliveries",
"export_jobs","saved_reports","report_schedules",
"feature_flags","audit_logs","security_events","outbox_events",
]
for t in tables:
    bullet(t)
page_break()

h1("61. Required Human-Readable Codes")
table(["Object","Format example","Purpose"], [
("Company","CMP-000127","Human support/search"),
("Site","SITE-BLR01","Site identity"),
("Process","PROC-P01","Process identity"),
("Team","TEAM-CC01","Team identity"),
("Employee","EMP-004582","Employee-facing operational reference"),
("Vendor","VND-00074","Vendor operations"),
("Driver","DRV-001892","Driver operations"),
("Vehicle","VEH-KA01AB1234","Fleet"),
("Vehicle Duty","DUT-20260905-0001","Duty identity"),
("QR","QR-20260905-0001","QR lifecycle"),
("Route","RTE-BLR-001","Route identity"),
("Schedule","SCH-20260905-003","Schedule instance"),
("Booking","BKG-20260905-0042","Booking"),
("Trip","TRP-20260905-001274","Trip"),
("Incident","INC-20260905-0004","Incident"),
("Document","DOC-DRV-001892-0003","Document"),
("Invoice","INV-202609-00129","Invoice"),
("Export","EXP-20260905-0021","Export job"),
("Case","CAS-20260905-0042","Support/operational case"),
])
para("Codes are unique within an appropriate business namespace. UUIDs remain the technical primary keys.")
page_break()

h1("62. API Design")
h2("62.1 Core endpoint families")
bl(["/api/v1/platform/owners", "/api/v1/platform/companies", "/api/v1/sites", "/api/v1/processes", "/api/v1/users", "/api/v1/roles", "/api/v1/permissions", "/api/v1/transport/employees", "/api/v1/bookings", "/api/v1/trips", "/api/v1/drivers", "/api/v1/vehicles", "/api/v1/vehicle-qr", "/api/v1/routes", "/api/v1/schedules", "/api/v1/gps/events", "/api/v1/control-room/live", "/api/v1/documents", "/api/v1/reports", "/api/v1/billing/internal", "/api/v1/billing/external", "/api/v1/audit"])
h2("62.2 Example trip APIs")
code("""POST  /api/v1/trips/:id/assign
POST  /api/v1/trips/:id/unassign
POST  /api/v1/trips/:id/reassign
POST  /api/v1/trips/:id/vehicle-change
POST  /api/v1/trips/:id/passengers
POST  /api/v1/trips/:id/passengers/reassign
POST  /api/v1/trips/:id/passengers/remove
POST  /api/v1/trips/:id/split
POST  /api/v1/trips/:id/merge
POST  /api/v1/trips/:id/route-change
POST  /api/v1/trips/:id/cancel
GET   /api/v1/trips/:id/history
""")
page_break()

h1("63. Idempotency and Concurrency")
para("Dispatch actions are race-sensitive. The server must prevent two operators from assigning the same driver or exceeding vehicle capacity due to simultaneous submissions.")
bl(["Idempotency keys for critical POST actions.", "Optimistic concurrency/version checks.", "Database transactions.", "Conflict response with current state.", "Do not trust stale browser state."])
page_break()

h1("64. RLS and Database Security")
para("Enable Postgres Row Level Security on exposed tenant-sensitive tables. Supabase currently documents RLS as database-level granular authorization and recommends enabling it for exposed tables, with tests for allow/deny behavior. citeturn433501search0")
bl(["Use RLS as defense in depth.", "Keep grants aligned with application needs.", "Test select/insert/update/delete allow/deny cases.", "Review views because views can bypass RLS if created with elevated ownership.", "Never expose service-role credentials to clients."])
page_break()

h1("65. Storage and Privacy")
bl(["Private buckets for sensitive documents.", "Short-lived access URLs.", "Download authorization.", "Retention policy.", "Deletion/archive policy.", "Access audit.", "Mask sensitive fields based on role."])
page_break()

h1("66. Search and Filtering Standards")
bl(["Company/site/process/date filters.", "Status filters.", "Vendor.", "Driver.", "Vehicle.", "Shift.", "Free-text search.", "Saved views.", "Clear filter chips.", "Reset filters.", "Pagination/virtualization."])
page_break()

h1("67. Bulk Operations")
bl(["Bulk employee import.", "Bulk employee update.", "Bulk offboard.", "Bulk booking.", "Bulk trip assignment where appropriate.", "Bulk passenger movement.", "Bulk document review.", "Bulk export."])
para("Every bulk action shows selected count, valid count, blocked count, reason per blocked record and a final result.")
page_break()

h1("68. Import and Data Quality")
bl(["CSV/Excel upload.", "Column mapping.", "Required-field checks.", "Duplicate detection.", "Reference-data validation.", "Coordinate validation.", "Preview.", "Import job.", "Success/failure summary.", "Downloadable error list.", "Audit."])
page_break()

h1("69. Scheduling Background Jobs")
bl(["Exports.", "Large imports.", "Notifications.", "Webhook delivery.", "Document expiry scans.", "Billing aggregation.", "Forecasting.", "Route optimization.", "GPS aggregation.", "Daily summaries."])
page_break()

h1("70. Forecasting and Planning")
h2("70.1 Forecast layers")
bl(["Long range.", "Monthly.", "Weekly.", "Day ahead.", "Last minute."])
h2("70.2 Inputs")
bl(["Historical bookings.", "Cancellations.", "No-shows.", "Shift patterns.", "Holiday calendars.", "Headcount.", "Occupancy.", "Site/process demand."])
h2("70.3 Outputs")
bl(["Expected passengers.", "Expected vehicles.", "Capacity risk.", "Likely shortages.", "Likely overcapacity.", "Recommended staffing."])
page_break()

h1("71. Cost and Sustainability Analytics")
bl(["Cost per trip.", "Cost per employee.", "Cost per km.", "Empty-seat cost.", "Vendor comparison.", "Vehicle utilization.", "Fuel/energy estimates.", "CO2 estimate where data supports it."])
page_break()

h1("72. Vendor Performance")
bl(["On-time rate.", "Cancellation rate.", "GPS availability.", "Incidents.", "Complaints.", "SLA attainment.", "Cost.", "Document compliance.", "Driver availability.", "Trip acceptance."])
page_break()

h1("73. Service-Level Management")
bl(["Driver arrival SLA.", "Pickup SLA.", "Trip completion SLA.", "Vendor response SLA.", "Incident response SLA.", "GPS availability SLA."])
para("SLA thresholds are versioned by contract or policy and can be reported historically.")
page_break()

h1("74. Compliance Automation")
bl(["Driver licence expiry.", "Vehicle registration/fitness expiry.", "Insurance expiry.", "Vendor contract expiry.", "Permit expiry.", "Training expiry.", "Document verification queue."])
para("Expired compliance documents can block dispatch when the applicable policy says so. The UI must clearly explain the blocker.")
page_break()

h1("75. Support and Operational Cases")
para("Enterprise customers need a support workflow for issues such as wrong passenger move, incorrect driver assignment, billing mismatch, document rejection and integration failure.")
bl(["Case Code.", "Category.", "Priority.", "Related company/site/process.", "Related trip/employee/driver/vehicle.", "Owner.", "Status.", "Comments.", "Attachments.", "Resolution.", "Audit."])
page_break()

h1("76. Audit-Friendly Timeline")
para("Every major object should have a chronological timeline showing business events. Users should be able to understand what happened without opening raw logs.")
bl(["Booking created.", "Driver assigned.", "Vehicle changed.", "Passenger moved.", "No-show.", "GPS anomaly.", "Document uploaded.", "Invoice generated.", "Permission changed.", "Export downloaded."])
page_break()

h1("77. Data Retention and Archival")
bl(["Audit retention.", "GPS retention.", "Trip retention.", "Document retention.", "Export retention.", "Notification retention.", "Invoice retention."])
para("Retention periods are configurable by company policy and applicable legal requirements. Do not hard-delete operational history as a side effect of ordinary edits.")
page_break()

h1("78. Observability and Operations")
bl(["API health.", "Database health.", "Queue health.", "GPS ingestion.", "Map provider health.", "Notification delivery.", "Webhook delivery.", "Export jobs.", "Authentication failures.", "Dispatch failures.", "Background job latency."])
page_break()

h1("79. Environment Strategy")
table(["Environment","Purpose"], [
("Local","Developer implementation and isolated data"),
("Development","Shared feature development"),
("Staging","Production-like validation"),
("Production","Live customer operations"),
])
bl(["Separate secrets.", "Separate map credentials where practical.", "No production demo users.", "Database migrations versioned.", "Seed data isolated.", "Rollback procedure."])
page_break()

h1("80. Secret and API Key Management")
bl(["Environment variables or managed secret store.", "Never commit secrets.", "Never place server keys in frontend bundles.", "Separate browser and server map credentials.", "Restrict Google keys by application and API.", "Monitor usage.", "Rotate carefully.", "Document ownership and expiry of secrets."])
para("Google's current security guidance explicitly recommends restricting keys, using separate keys for each application, splitting client-side and server-side use, and disabling unused services. citeturn749546search0")
page_break()

h1("81. Production Performance")
bl(["Pagination.", "Virtualized large tables.", "Indexed company/site/process keys.", "Async exports.", "Cached reference data.", "Efficient live-map updates.", "Batch GPS writes.", "Background billing aggregation.", "Read models for high-volume dashboards."])
page_break()

h1("82. Read Models")
para("Operational dashboards and reports should use optimized read models/projections where necessary rather than forcing every screen to scan high-write transactional tables.")
bl(["Control-room projection.", "Driver availability projection.", "Vehicle status projection.", "Trip summary projection.", "Report datasets.", "Executive KPI projection."])
page_break()

h1("83. Outbox and Event Model")
bl(["company.created", "user.role_changed", "user.permission_changed", "employee.onboarded", "employee.offboarded", "booking.created", "booking.cancelled", "trip.assigned", "trip.reassigned", "trip.passenger_moved", "driver.state_changed", "vehicle_qr.created", "vehicle_qr.scanned", "gps.received", "document.expired", "incident.created", "invoice.created", "export.completed"])
para("Transactional state changes write an outbox event in the same database transaction. Workers consume those events for notifications, reporting projections and integrations.")
page_break()

h1("84. Notification Templates")
table(["Template","Audience","Example"], [
("TRIP_ASSIGNED","Employee/driver","Your vehicle and driver have been assigned."),
("TRIP_DELAYED","Employee/manager","Your pickup is delayed."),
("PASSENGER_MOVED","Employee","Your trip has been changed."),
("DOCUMENT_EXPIRING","Driver/vendor","Your licence expires soon."),
("GPS_STALE","Transport Admin","Vehicle GPS has been stale for X minutes."),
("REPORT_READY","Requester","Your report is ready to download."),
("INVOICE_READY","Finance","Invoice is ready for review."),
])
page_break()

h1("85. Authorization Test Matrix")
tests = [
("Super Admin creates company","Allowed","Company created + audit"),
("Manager creates company","Denied","403 no mutation"),
("Transport Admin onboards employee","Allowed","Employee persisted"),
("Coordinator onboards employee","Denied","403"),
("Transport Admin generates QR","Allowed","QR created"),
("Manager generates QR","Denied","403"),
("Coordinator generates QR without delegation","Denied","403"),
("Coordinator generates QR with delegation","Allowed","Within scope"),
("Transport Admin moves one passenger","Allowed","Only selected passenger changes"),
("Transport Admin moves passenger to incompatible trip","Denied","Reason shown"),
("Employee accesses another employee","Denied","No data leak"),
("Vendor A accesses Vendor B","Denied","403"),
("Sleepy driver assigned","Denied","Compliance blocker"),
("Expired licence driver assigned","Denied","Document blocker"),
("Permission OFF then API called","Denied","403"),
("Role changed then old API called","Denied","403"),
("Unauthorized report export","Denied","403"),
("Authorized report export","Allowed","Scoped XLSX/PDF/CSV"),
]
table(["Scenario","Expected","Acceptance"], tests)
page_break()

h1("86. End-to-End Acceptance Journey")
para("A production tester must be able to execute the following without direct database editing.")
journey = [
"SAAS_OWNER creates MOVEINSYNC_OWNER.",
"MOVEINSYNC_OWNER creates SUPER_ADMIN.",
"SUPER_ADMIN creates customer company.",
"SUPER_ADMIN creates sites and processes.",
"SUPER_ADMIN creates customer admin and Transport Admin.",
"Transport Admin signs in.",
"Transport Admin onboards employee.",
"Transport Admin adds/validates pickup point.",
"Transport Admin creates driver and vehicle.",
"Transport Admin uploads compliance documents.",
"Transport Admin creates vehicle duty.",
"Transport Admin generates vehicle-shift QR.",
"Employee receives account activation flow.",
"Employee books transport.",
"Trip is generated with Trip Code.",
"Transport Admin assigns driver and vehicle.",
"Driver becomes AVAILABLE.",
"Driver app receives duty/trip.",
"Driver activates duty / uses QR.",
"GPS becomes LIVE.",
"Trip starts with configured OTP/arrival rules.",
"One passenger becomes NO_SHOW.",
"Transport Admin moves another passenger to a compatible trip.",
"Transport Admin unassigns the driver.",
"Transport Admin assigns a replacement driver.",
"Trip continues with updated assignment history.",
"Trip completes.",
"Operational charge is generated.",
"Invoice is generated.",
"Finance reconciles.",
"Transport Admin downloads operational Excel/PDF.",
"Super Admin downloads cross-company report.",
"Audit timeline shows every sensitive change."
]
for i, step in enumerate(journey,1):
    numbered(f"{i}. {step}")
page_break()

h1("87. Real-Life Operational Edge Cases")
bl([
    "Driver calls in sick 10 minutes before pickup.",
    "Vehicle breaks down after dispatch.",
    "One passenger asks to move to another cab.",
    "One passenger is a no-show while others board.",
    "A new passenger must be added to an existing cab.",
    "Route becomes unavailable because of traffic or road closure.",
    "Female transport safety rule makes an apparently optimal cab invalid.",
    "Driver document expires before a scheduled trip.",
    "GPS stops sending data during an active trip.",
    "Mobile network disappears and driver reconnects later.",
    "Two coordinators attempt to assign the same driver.",
    "Trip crosses midnight.",
    "Employee is offboarded after booking but before pickup.",
    "Vendor invoice includes a trip not present in operational history.",
    "Report export is too large for synchronous generation.",
    "Document is rejected and needs re-upload.",
    "Customer changes billing level from company to process.",
    "A site is temporarily suspended but historical reporting must remain available.",
])
page_break()

h1("88. Trip Operations Decision Matrix")
table(["Situation","Operator action","System behavior"], [
("Driver unavailable before dispatch","Change Driver","Revalidate driver + notify"),
("Vehicle unavailable before dispatch","Change Vehicle","Revalidate capacity + documents"),
("One passenger must change cab","Move Passenger","Revalidate destination + create source/destination events"),
("Passenger no-show","Mark No-show","Release capacity if policy permits"),
("Route unsafe","Change Route","Safety validation before commit"),
("Trip started","Attempt passenger move","Block or require emergency override"),
("Trip completed","Attempt reassignment","Block; preserve history"),
("Duplicate assignment","Second operator assigns same driver","409/conflict and refresh current state"),
("QR expired","Driver scans","Reject and show replacement path"),
])
page_break()

h1("89. UI Acceptance - Transport Admin")
bl(["Can find a trip by Trip Code.", "Can see driver, vehicle and passengers immediately.", "Can remove one passenger without cancelling the whole trip.", "Can move a passenger to another eligible trip.", "Can change driver.", "Can change vehicle.", "Can see operational impact before confirming.", "Can generate QR.", "Can see documents.", "Can upload documents.", "Can download reports.", "Can understand errors without technical knowledge."])
page_break()

h1("90. UI Acceptance - Employee")
bl(["Can understand today's trip on first screen.", "Can book own transport.", "Can see pickup.", "Can see ETA.", "Can cancel permitted booking.", "Can view history.", "Can report issue.", "Can access safety action.", "Does not see administrative modules."])
page_break()

h1("91. UI Acceptance - Driver")
bl(["Can see current duty.", "Can become active.", "Can see next passenger.", "Can mark arrived.", "Can mark no-show.", "Can start/complete trip.", "Can mark break/sleepy.", "Can trigger emergency.", "Can see GPS/network state."])
page_break()

h1("92. UI Acceptance - Finance")
bl(["Can distinguish internal SaaS billing from external transport billing.", "Can see relevant invoices.", "Can reconcile.", "Can export finance reports.", "Cannot accidentally dispatch a vehicle unless separately authorized."])
page_break()

h1("93. UI Acceptance - Super Admin")
bl(["Can access every functional module.", "Can create/manage companies.", "Can create/manage owners according to hierarchy.", "Can manage roles and permissions.", "Can perform operational trip changes.", "Can generate/revoke QR.", "Can upload/download documents.", "Can access all reports.", "Can download all authorized report formats.", "Can access both billing domains.", "Can use Access Simulator.", "Can inspect system health and audit."])
page_break()

h1("94. Report Permission Matrix")
table(["Role","Operational reports","Finance reports","Audit reports","Cross-company","Excel/PDF/CSV"], [
("SAAS_OWNER","All","All","All","Yes","All"),
("MOVEINSYNC_OWNER","All","All","All","Yes","All"),
("SUPER_ADMIN","All","All","All","Yes","All"),
("TRANSPORT_HEAD","All in customer scope","Cost as allowed","Limited","No unless granted","All allowed"),
("TRANSPORT_ADMIN","Operational","Cost if granted","Limited","Customer scope","Per permission"),
("TRANSPORT_COORDINATOR","Dispatch","No unless granted","No","No","Operational only"),
("MANAGER","Team","No","No","No","Team scope"),
("EMPLOYEE","Own","No","No","No","Own where offered"),
("FINANCE_ADMIN","Finance","All company finance","No","No","All authorized finance"),
("FINANCE_VIEWER","Finance read-only","Read-only","No","No","Authorized"),
("VENDOR_FINANCE","Vendor finance","Vendor only","No","No","Authorized vendor"),
("AUDITOR","Audit/compliance","As granted","All audit","Per assigned audit scope","Authorized"),
])
page_break()

h1("95. Billing Permission Matrix")
table(["Permission group","Who normally uses it"], [
("Internal SaaS billing","SAAS_OWNER, MOVEINSYNC_OWNER, SUPER_ADMIN, PLATFORM_FINANCE_*"),
("External rate cards","SUPER_ADMIN, TRANSPORT_HEAD where granted, FINANCE_ADMIN, authorized vendor finance"),
("External invoice review","FINANCE_ADMIN, FINANCE_APPROVER, SUPER_ADMIN"),
("External reconciliation","FINANCE_ADMIN, SUPER_ADMIN, authorized finance analyst"),
("Finance exports","Finance roles + SUPER_ADMIN/owners"),
])
page_break()

h1("96. Document Permission Matrix")
table(["Role","View","Upload","Verify","Reject","Download"], [
("SUPER_ADMIN / Owners","All","All","All","All","All"),
("TRANSPORT_ADMIN","Assigned scope","Drivers/vehicles/employees where applicable","Assigned scope","Assigned scope","Assigned scope"),
("VENDOR_COMPLIANCE_MANAGER","Vendor scope","Vendor docs","Vendor docs","Vendor docs","Vendor docs"),
("DRIVER","Own","Own required docs","No","No","Own docs"),
("EMPLOYEE","Own permitted docs","Own permitted docs","No","No","Own docs"),
("FINANCE","Finance docs","Invoice attachments where authorized","No unless granted","No unless granted","Finance scope"),
])
page_break()

h1("97. Production Hardening Checklist")
bl(["Versioned migrations.", "Backups and restore tests.", "Secure authentication.", "MFA for privileged roles.", "Rate limiting.", "Input/schema validation.", "Tenant isolation tests.", "Authorization tests.", "Structured logs.", "Correlation IDs.", "Health/readiness endpoints.", "Queue retry/dead-letter strategy.", "GPS offline handling.", "Export isolation.", "Monitoring/alerts.", "No production demo accounts unless explicitly enabled.", "CI tests.", "Deployment migration/rollback procedure."])
page_break()

h1("98. Release Gates")
table(["Gate","Must be true"], [
("Identity","Authentication, recovery and MFA policies work"),
("Authorization","Role/permission/scope tests pass"),
("Tenant isolation","Cross-tenant access tests pass"),
("Transport","Employee/booking/trip workflows complete"),
("Dispatch","Driver/vehicle/passenger modifications work"),
("QR","Generation/use/revocation works"),
("GPS","Ingestion and stale/offline states work"),
("Documents","Upload/verify/reject/expiry work"),
("Billing","Internal and external billing separated"),
("Reports","Scoped Excel/PDF/CSV downloads work"),
("Audit","Sensitive mutations traceable"),
("UX","No fake states; errors and empty states are usable"),
("Operations","Health, jobs, alerts and restore tested"),
])
page_break()

h1("99. Mimo Build Order")
phases = [
("Phase 1","Supabase + Postgres + RLS + Auth + base schema"),
("Phase 2","Owner hierarchy + RBAC + permission engine + scopes"),
("Phase 3","Companies + sites + processes + teams"),
("Phase 4","Users + invitations + employee transport master"),
("Phase 5","Pickup points + maps + geocoding/validation"),
("Phase 6","Vendors + drivers + vehicles + documents"),
("Phase 7","Vehicle duty + QR"),
("Phase 8","Bookings + shifts + roster + schedules"),
("Phase 9","Trips + dispatch + passenger operations"),
("Phase 10","GPS + driver app + employee app + control room"),
("Phase 11","Safety + incidents + guard"),
("Phase 12","Billing + rate cards + invoices + reconciliation"),
("Phase 13","Reports + exports + dashboards"),
("Phase 14","Integrations + webhooks + APIs"),
("Phase 15","Forecasting + optimization + intelligence"),
("Phase 16","Hardening + performance + acceptance + production release"),
]
table(["Phase","Deliverables"], phases)
page_break()

h1("100. Mimo Completion Protocol")
bl([
    "Inspect the existing code before changing it.",
    "Do not duplicate business rules in multiple components.",
    "Do not create fake endpoints.",
    "Do not leave placeholders for required production workflows.",
    "Complete database and API before wiring the final UI action.",
    "Use the acceptance tests as implementation checkpoints.",
    "After each phase, run build/lint/tests and verify the UI.",
    "When a rule is unclear, preserve the business authority model rather than inventing a broader permission.",
    "Keep role, permission and scope separate.",
    "Preserve history whenever operational state changes.",
    "Do not reduce functionality to fit a page count.",
])
page_break()

h1("101. Master Non-Negotiable Rules")
rules = [
"SUPER_ADMIN is full-functional owner-level access.",
"SAAS_OWNER is the highest governance identity.",
"MOVEINSYNC_OWNER has complete functional access and owner governance below SAAS_OWNER.",
"Super Admin can perform normal operational transport actions including driver reassignment, vehicle replacement and passenger movement.",
"Transport Admin owns transport employee onboarding/offboarding.",
"Managers do not automatically receive QR generation.",
"Finance does not automatically receive dispatch or QR privileges.",
"Vehicle QR generation belongs to Transport Admin by default and can be delegated explicitly.",
"Trip Code is mandatory for every trip.",
"Passenger movement is an independent workflow.",
"Trip history is never silently overwritten.",
"Documents are a real module with upload and lifecycle.",
"Report download is a permissioned capability.",
"Internal SaaS billing and external transport billing are separate domains.",
"Company, site, process and team are normalized scope dimensions.",
"Server-side authorization is mandatory.",
"Supabase RLS provides database-level defense in depth.",
"Maps use a provider abstraction and restricted credentials.",
"Safety outranks optimization.",
"Page count is never a reason to remove functionality.",
"UI language must be understandable to non-technical users.",
"Production completion requires end-to-end acceptance, not screenshots."
]
for r in rules:
    bullet(r)
page_break()

h1("102. Final Definition of Done")
para("Move In Sync is complete only when a clean environment supports the complete owner-to-operation journey, every critical permission is enforced server-side, every required operational state is persisted, every sensitive action is auditable, every report export is scoped, every document workflow is functional, and every primary user can complete their work from an understandable interface.")
para("The product is not considered production ready if a screen merely renders, if metrics are hard-coded, if a button does nothing, if an API accepts unauthorized scope, if an export leaks data, or if a critical operation requires direct database editing.")
para("The specification intentionally contains more screens and more detail than a prototype brief. That is deliberate. The platform must be built for real enterprise operations, not page-count optimization.")
page_break()

h1("Appendix A - Suggested Environment Variables")
code("""# Supabase
NEXT_PUBLIC_SUPABASE_URL=
NEXT_PUBLIC_SUPABASE_ANON_KEY=
SUPABASE_SERVICE_ROLE_KEY=

# Google Maps browser
NEXT_PUBLIC_GOOGLE_MAPS_BROWSER_KEY=

# Google Maps server
GOOGLE_MAPS_GEOCODING_KEY=
GOOGLE_MAPS_ADDRESS_VALIDATION_KEY=
GOOGLE_MAPS_ROUTES_KEY=
GOOGLE_MAPS_ROADS_KEY=
GOOGLE_MAPS_STATIC_KEY=

# Route Optimization
GOOGLE_ROUTE_OPTIMIZATION_SERVICE_ACCOUNT=

# Notifications
EMAIL_PROVIDER_API_KEY=
SMS_PROVIDER_API_KEY=
PUSH_PROVIDER_KEY=
WHATSAPP_PROVIDER_API_KEY=

# Integrations
HRMS_API_URL=
HRMS_API_KEY=
ERP_API_URL=
ERP_API_KEY=

# App
APP_URL=
API_URL=
ENVIRONMENT=
""")
para("Only the browser key belongs in a client bundle, and even that key must be restricted to the production website/application and only the APIs required by the client. Server-side credentials remain secret.")
page_break()

h1("Appendix B - Example Trip Detail UX")
code("""TRIP TRP-20260905-001274
Status: DISPATCHED

Driver: Ravi Kumar       Vehicle: KA01AB1234
Vendor: Vendor Alpha    Shift: 18:30 - 03:30
Site: BLR01             Process: P01

PASSENGERS (4)
[ ] EMP-001  Rahul       Assigned
[ ] EMP-002  Priya       Assigned
[ ] EMP-003  Amit        No-show
[ ] EMP-004  Neha        Assigned

PRIMARY ACTIONS
[ Change Driver ] [ Change Vehicle ] [ Move Passenger ]
[ Add Passenger ] [ Remove Passenger ] [ Change Route ]

MAP
- current vehicle
- pickup sequence
- next stop
- ETA

TIMELINE
18:04 Driver assigned
18:10 Vehicle assigned
18:16 QR generated
18:29 Driver arrived
18:31 Passenger EMP-003 marked no-show
18:32 EMP-004 moved to TRP-20260905-001301
""")
page_break()

h1("Appendix C - Example Passenger Move Workflow")
code("""1. Select passenger EMP-004
2. Click Move Passenger
3. Search or browse eligible destination trips
4. System checks:
   - capacity
   - time window
   - route compatibility
   - site/process rules
   - female safety
   - destination trip state
   - driver/vehicle capacity
5. Show impact
6. Ask reason
7. Confirm
8. Transaction updates source + destination
9. Write events
10. Notify affected passenger/operator
11. Refresh trip cards and control-room map
""")
page_break()

h1("Appendix D - Example Document Workflow")
code("""Vehicle VEH-KA01AB1234
Documents
  Insurance.pdf       VERIFIED
  Registration.pdf    VERIFIED
  Fitness.pdf         EXPIRING SOON

[ Upload Document ]

Upload dialog
  Document type
  Document number
  Issue date
  Expiry date
  File
  [Upload]

Result:
  Status = PENDING_REVIEW

Reviewer:
  [Verify] [Reject / Request Re-upload]

Audit:
  Uploaded by ...
  Verified by ...
""")
page_break()

h1("Appendix E - Sources Used for Current Provider Guidance")
para("The functional requirements in this document are based primarily on the supplied Move In Sync V8 material and the later user-requested expansions. Current provider-specific notes in the map/authentication sections were checked against current official Google Maps Platform and Supabase documentation.")
para("Google Maps Platform security guidance: restricted keys, separate application keys, client/server separation, API restrictions and disabling unused services. citeturn749546search0")
para("Google Routes API: Compute Route Matrix supports up to 625 route elements. citeturn433501search2")
para("Google Address Validation API: address validation, geocoding information and deliverability-oriented results. citeturn433501search1turn433501search3")
para("Supabase Row Level Security: database-level granular authorization and guidance to enable RLS on exposed tables and test allow/deny paths. citeturn433501search0")

# Add a final index-like section with implementation checklists
page_break()
h1("Appendix F - Builder Checklist")
check_groups = {
"Identity": ["Supabase Auth configured","Invitation/recovery works","MFA policy works","Sessions can be revoked"],
"RBAC": ["Owner hierarchy seeded","All roles seeded","Permission catalogue seeded","ON/OFF overrides work","Scope assignments work","Access simulator works"],
"Transport": ["Employee master","Pickup points","Bookings","Trips","Dispatch","Passenger movement","Driver/vehicle reassignment"],
"Fleet": ["Vendor","Driver","Vehicle","Documents","Maintenance","Duty","QR"],
"Safety": ["SOS","Female transport","Incidents","No-show","Driver fatigue","Guard verification"],
"Maps/GPS": ["Map provider abstraction","Restricted keys","Geocoding","Address validation","Routes","Matrix","Live GPS","Stale/offline"],
"Billing": ["Internal SaaS","External transport","Rate cards","Invoices","Reconciliation"],
"Reports": ["Catalogue","Builder","Saved reports","Schedules","Excel","PDF","CSV","Audit"],
"Platform": ["Integrations","Webhooks","Feature flags","Health","Jobs","Audit"],
"UX": ["Role home","Friendly navigation","Responsive UI","Accessibility","Loading/error/empty/403 states"],
"Production": ["Migrations","Backups","Restore","Monitoring","Rate limits","Authorization tests","Tenant isolation tests"],
}
for group, items in check_groups.items():
    h2(group)
    for item in items:
        bullet("[ ] " + item)

doc.save(out)
print(out)
